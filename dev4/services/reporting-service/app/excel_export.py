"""Export Excel colorie par statut avec openpyxl."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Couleurs par statut
FILL_NORMAL = PatternFill("solid", fgColor="C6EFCE")  # vert
FILL_HIGH   = PatternFill("solid", fgColor="FFC7CE")  # rouge
FILL_LOW    = PatternFill("solid", fgColor="BDD7EE")  # bleu
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")  # bleu fonce
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILL = {"NORMAL": FILL_NORMAL, "HIGH": FILL_HIGH, "LOW": FILL_LOW}


def build_excel(analyses: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Analyses"

    headers = ["Analyse", "Valeur", "Unite", "Ref min", "Ref max",
               "Statut", "Date"]
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER

    for a in analyses:
        ws.append([
            a.analysis_name, a.value, a.unit or "",
            a.ref_min, a.ref_max, a.status,
            a.analysis_date.isoformat() if a.analysis_date else "",
        ])
        row = ws.max_row
        fill = STATUS_FILL.get(a.status)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=row, column=col)
            c.border = BORDER
            if fill:
                c.fill = fill

    # Largeur des colonnes
    widths = [22, 10, 12, 10, 10, 10, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
